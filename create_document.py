from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

out = Path('Ecomm_Inventory_Movement_Technical_Documentation.docx')
doc = Document()
sec = doc.sections[0]
sec.top_margin=Inches(.68); sec.bottom_margin=Inches(.68); sec.left_margin=Inches(.78); sec.right_margin=Inches(.78)
styles=doc.styles
styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(9.5); styles['Normal'].font.color.rgb=RGBColor(45,55,65)
styles['Normal'].paragraph_format.space_after=Pt(5)
for nm,size,color in [('Title',28,'17324D'),('Heading 1',18,'17324D'),('Heading 2',12,'187C80'),('Heading 3',10,'17324D')]:
 s=styles[nm]; s.font.name='Aptos Display' if nm!='Heading 3' else 'Aptos'; s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor.from_string(color); s.paragraph_format.space_before=Pt(10); s.paragraph_format.space_after=Pt(4)
styles['Title'].paragraph_format.space_after=Pt(8)
# header/footer
hp=sec.header.paragraphs[0]; hp.text='ECOMM INVENTORY MOVEMENT  /  TECHNICAL DOCUMENTATION'; hp.style='Caption'; hp.runs[0].font.color.rgb=RGBColor(24,124,128); hp.runs[0].font.bold=True
fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT; rr=fp.add_run('Engineering reference  |  20 October 2026'); rr.font.size=Pt(8); rr.font.color.rgb=RGBColor(110,120,128)

def shade(cell, fill):
 tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)
def table(headers, rows, widths=None):
 t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Light Shading Accent 1'
 for i,h in enumerate(headers):
  c=t.rows[0].cells[i]; c.text=h; shade(c,'17324D')
  for r in c.paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(8.5)
 for row in rows:
  cells=t.add_row().cells
  for i,v in enumerate(row):
   cells[i].text=str(v); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   for p in cells[i].paragraphs:
    p.paragraph_format.space_after=Pt(2)
    for r in p.runs: r.font.size=Pt(8.2)
 if widths:
  for row in t.rows:
   for i,w in enumerate(widths): row.cells[i].width=Inches(w)
 doc.add_paragraph().paragraph_format.space_after=Pt(1)
 return t
def para(txt='', style=None): return doc.add_paragraph(txt, style)
def bullet(txt): doc.add_paragraph(txt,style='List Bullet')
def code(txt):
 p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.15); p.paragraph_format.space_before=Pt(3); p.paragraph_format.space_after=Pt(6)
 r=p.add_run(txt); r.font.name='Consolas'; r.font.size=Pt(8); r.font.color.rgb=RGBColor(30,70,85)
 pPr=p._p.get_or_add_pPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),'EEF3F5'); pPr.append(shd)

doc.add_paragraph('Ecomm Inventory Movement',style='Title')
p=doc.add_paragraph('Technical documentation and server deployment guide'); p.runs[0].font.size=Pt(14); p.runs[0].font.color.rgb=RGBColor(24,124,128)
para('This guide describes the current Spring Boot application, the Oracle and inventory-service dependencies it needs, local setup, and a repeatable Linux server deployment. The application is a server-rendered web UI with JSON endpoints; inventory reads come from Oracle, while adjustment operations call an external OAuth-protected inventory API. The service must have network access and credentials for both systems before it can perform useful work.')
table(['Document details','Value'],[['Application','Ecomm Inventory Movement'],['Build','Maven wrapper; Spring Boot 4.1.1; Java 17'],['Persistence','Spring Data JPA with Oracle JDBC'],['Web','Spring MVC, Thymeleaf, Bean Validation'],['External integration','Spring RestClient and OpenFeign dependency; adjustment flow uses RestClient'],['Deployment pattern in this guide','Linux systemd service behind a TLS reverse proxy']])

doc.add_heading('1. System overview',level=1)
para('The application starts from EcommInventoryMovementApplication and enables Feign client scanning. InventoryController serves Thymeleaf pages and JSON routes. InventoryService reads the INVENTORY_LI table through Spring Data JPA. InventoryAdjustmentService obtains an OAuth client-credentials token, queries the inventory viewer API, and submits adjustment uploads. The browser never receives the API credentials or bearer token.')
table(['Layer','Current responsibility'],[['Web interface','Thymeleaf pages for inventory movement and job adjustment; browser JavaScript calls JSON endpoints.'],['Controller','Validates requests and delegates inventory reads or adjustment work.'],['Service','InventoryService reads Oracle; InventoryAdjustmentService invokes OAuth, viewer, and upload APIs.'],['Repository and entity','InventoryRepository maps NKW.INVENTORY_LI; InventoryLi uses WHOUSE, STYLE, COLOR_NO as its composite key.'],['External systems','Oracle database and configured inventory adjustment API URLs.']])
para('No Dockerfile, compose file, database migration, scheduler configuration, or production reverse-proxy configuration is present in the repository. The server procedure below therefore uses the built executable JAR and systemd. “Scheduled adjustment” currently means a request to an HTTP endpoint; the code does not create a recurring timer or background schedule.',style=None)

doc.add_heading('2. Current application behavior',level=1)
table(['Method and route','Purpose'],[['GET /','Controller mapping collision risk; see deployment notes.'],['GET /job-inventory-adjustment','Renders the manual and adjustment page.'],['GET /plm-styles-integration','Renders the PLM integration page and loads inventory.'],['GET /inventory-movement','Renders inventory movement page and loads inventory.'],['GET /inventory','Returns all mapped inventory rows as JSON.'],['GET /inventory/style/{style}','Returns inventory rows matching a style.'],['POST /inventory/adjustment/manual','Posts one warehouse/style/color/size/quantity adjustment.'],['POST /inventory/adjustment/beta','Validates stock at DC and CA, then submits +5 to DC and -5 to CA.'],['POST /inventory/adjustment/scheduled','Scans sizes 1–18 using provided CA/DC stock thresholds and transfer quantity.']])
para('The manual request requires nonblank warehouse, style, color, and size plus a positive quantity. The beta request requires nonblank style and color and positive size. The scheduled request requires nonnegative minimum/maximum thresholds, ordered min/max pairs, and a positive transfer quantity. Scheduled processing returns movement and failure counts; it does not run automatically.')
para('Operations note: InventoryController currently declares “/” in three separate page mappings. This creates ambiguous handler registration for the root path and can prevent startup or cause unpredictable root behavior. Remove “/” from all but one mapping, or define one deliberate landing-page mapping, before production release.')

doc.add_heading('3. Prerequisites and access',level=1)
for x in ['JDK 17 on developer machines and the server. Confirm with java -version.','Maven Wrapper is included (mvnw on Unix, mvnw.cmd on Windows); Maven downloads dependencies during the first build.','An Oracle account with SELECT access to NKW.INVENTORY_LI and connectivity from the app host. The mapped table and columns must match the entity.','Inventory API URLs, OAuth client username/password, and outbound HTTPS access from the host. Use environment-specific credentials and approved endpoints.','A supported Linux host such as Ubuntu Server or another systemd distribution, a dedicated unprivileged service account, and a TLS reverse proxy for browser access.','Firewall rules that allow inbound proxy traffic and outbound Oracle/API traffic; keep the application port private to the host or trusted network.']:
 bullet(x)

doc.add_heading('4. Local setup',level=1)
para('Clone or copy the repository, then configure database and API settings outside source control. The checked-in application.yaml currently contains environment-specific values, including credentials. Rotate any real credentials that have been committed, remove them from the tracked configuration, and supply replacements through a secret manager or environment variables before building a production release.')
para('Spring Boot relaxed binding supports the following environment-variable names. The integration URLs are configurable as well as credentials:')
table(['Property','Environment variable','Notes'],[['spring.datasource.url','SPRING_DATASOURCE_URL','Oracle JDBC URL, for example jdbc:oracle:thin:@//dbhost:1521/service.'],['spring.datasource.username','SPRING_DATASOURCE_USERNAME','Least-privilege Oracle account.'],['spring.datasource.password','SPRING_DATASOURCE_PASSWORD','Inject through secret storage.'],['inventory.adjustment.oauth-url','INVENTORY_ADJUSTMENT_OAUTH_URL','OAuth token endpoint.'],['inventory.adjustment.upload-url','INVENTORY_ADJUSTMENT_UPLOAD_URL','Adjustment upload endpoint.'],['inventory.adjustment.viewer-url','INVENTORY_ADJUSTMENT_VIEWER_URL','Inventory viewer endpoint.'],['inventory.adjustment.oauth-username','INVENTORY_ADJUSTMENT_OAUTH_USERNAME','OAuth client ID/username.'],['inventory.adjustment.oauth-password','INVENTORY_ADJUSTMENT_OAUTH_PASSWORD','OAuth client secret/password.']])
para('The current YAML sets JPA schema management to none, which means the application will not create or update the schema. It also enables SQL logging. For production, keep schema management disabled, have the database owner provision the expected table, and set spring.jpa.show-sql=false to avoid verbose SQL output.')
code('''# Windows PowerShell example (use non-production credentials locally)\n$env:SPRING_DATASOURCE_URL = "jdbc:oracle:thin:@//localhost:1521/XEPDB1"\n$env:SPRING_DATASOURCE_USERNAME = "app_user"\n$env:SPRING_DATASOURCE_PASSWORD = "<local-secret>"\n$env:INVENTORY_ADJUSTMENT_OAUTH_USERNAME = "<client-id>"\n$env:INVENTORY_ADJUSTMENT_OAUTH_PASSWORD = "<client-secret>"\n.\\mvnw.cmd spring-boot:run''')
para('For a local Oracle connection, ensure the Oracle service name, listener port, account, and grants are valid. The project includes an OracleTest utility that uses TNS_ADMIN and an alias, but the application datasource in application.yaml defaults to a direct localhost service URL; these are separate connection examples.')

doc.add_heading('5. Build a release',level=1)
code('''# From the repository root\n./mvnw clean package\n# Windows PowerShell: .\\mvnw.cmd clean package\n# Result: target/ecomm-inventory-movement-0.0.1-SNAPSHOT.jar''')
para('The Spring Boot Maven plugin packages an executable JAR. Use a release pipeline to build from a reviewed commit, retain the build artifact, and promote the same artifact through test and production. Do not add secrets to the JAR, command-line arguments, or checked-in application.yaml. The project test dependencies are present, but this guide does not represent a test or build execution result.')

doc.add_heading('6. Linux server deployment',level=1)
doc.add_heading('6.1 Prepare the host',level=2)
code('''sudo useradd --system --home /opt/ecomm-inventory-movement --shell /usr/sbin/nologin ecomminv\nsudo install -d -o ecomminv -g ecomminv /opt/ecomm-inventory-movement/app\nsudo install -d -o root -g ecomminv -m 0750 /etc/ecomm-inventory-movement''')
para('Install a supported Java 17 runtime using the operating system package source approved by your organization. Confirm that the Oracle host and API endpoints resolve and are reachable from the server. Have the DBA create or approve the application account and grants before startup.')
doc.add_heading('6.2 Install the artifact and environment file',level=2)
code('''sudo install -o ecomminv -g ecomminv -m 0640 target/ecomm-inventory-movement-0.0.1-SNAPSHOT.jar /opt/ecomm-inventory-movement/app/app.jar\nsudoedit /etc/ecomm-inventory-movement/app.env''')
para('Store these values in app.env; replace examples with values from the target environment. Restrict file access to root and the service group, and use a secrets manager or systemd credential mechanism where available.')
code('''SPRING_PROFILES_ACTIVE=production\nSPRING_DATASOURCE_URL=jdbc:oracle:thin:@//oracle-host:1521/PROD\nSPRING_DATASOURCE_USERNAME=<service-account>\nSPRING_DATASOURCE_PASSWORD=<secret>\nSPRING_JPA_SHOW_SQL=false\nINVENTORY_ADJUSTMENT_OAUTH_URL=https://api.example/oauth/token\nINVENTORY_ADJUSTMENT_UPLOAD_URL=https://api.example/uploads/upload/INV_ADJ\nINVENTORY_ADJUSTMENT_VIEWER_URL=https://api.example/viewers/view/VR_INVENT_WH_Z\nINVENTORY_ADJUSTMENT_OAUTH_USERNAME=<client-id>\nINVENTORY_ADJUSTMENT_OAUTH_PASSWORD=<secret>''')
code('sudo chown root:ecomminv /etc/ecomm-inventory-movement/app.env && sudo chmod 0640 /etc/ecomm-inventory-movement/app.env')
para('A production profile is not currently defined in the repository; SPRING_PROFILES_ACTIVE=production only selects one if you add application-production.yaml or profile-specific settings. Environment variables shown above override the corresponding YAML properties directly.')
doc.add_heading('6.3 Configure systemd',level=2)
code('''# /etc/systemd/system/ecomm-inventory-movement.service\n[Unit]\nDescription=Ecomm Inventory Movement\nAfter=network-online.target\nWants=network-online.target\n\n[Service]\nType=simple\nUser=ecomminv\nGroup=ecomminv\nWorkingDirectory=/opt/ecomm-inventory-movement\nEnvironmentFile=/etc/ecomm-inventory-movement/app.env\nExecStart=/usr/bin/java -XX:MaxRAMPercentage=75 -jar /opt/ecomm-inventory-movement/app/app.jar\nSuccessExitStatus=143\nRestart=on-failure\nRestartSec=5\nNoNewPrivileges=true\nPrivateTmp=true\n\n[Install]\nWantedBy=multi-user.target''')
code('''sudo systemctl daemon-reload\nsudo systemctl enable --now ecomm-inventory-movement\nsudo systemctl status ecomm-inventory-movement\nsudo journalctl -u ecomm-inventory-movement -f''')
para('Spring Boot listens on port 8080 by default. Keep it bound behind an internal reverse proxy; configure the proxy for HTTPS, access logging, and an allowlist or SSO layer appropriate to the business. The application does not currently configure authentication or authorization for its pages or adjustment POST routes. Do not expose it publicly until an approved access-control layer is in place.')

doc.add_heading('7. Validation and operations',level=1)
para('After deployment, verify the service is active, startup logs show a successful datasource connection, and an authorized user can load the intended page. Check the Actuator dependency and the configured management exposure before relying on a health URL: no actuator exposure settings are currently present in application.yaml, so do not assume /actuator/health is reachable. Expose only the health details required by the infrastructure team and restrict management endpoints to trusted monitoring networks.')
table(['Operational area','Guidance'],[['Logs','Read systemd journal; alert on startup failures, database connection errors, OAuth failures, upload rejections, and viewer response errors. Avoid logging secrets or bearer tokens.'],['Database','Monitor connection capacity and query latency. Confirm schema/table grants and that Hibernate ddl-auto remains none.'],['External API','Monitor endpoint availability, certificate validity, token issuance, rate limits, and upload status.'],['Backups and recovery','Back up Oracle data according to database owner policy; preserve release artifact and environment configuration securely. Roll back by installing the last known-good JAR and restarting the service.'],['Updates','Build from reviewed source, stage in a non-production environment, then promote the same artifact. Restart systemd after replacing the JAR.'],['Access control','Put authenticated authorization at the application or trusted gateway before exposing adjustment actions. Protect state-changing requests with the organization’s CSRF and session-security design.']])

doc.add_heading('8. Production readiness items',level=1)
for x in ['Remove and rotate any credentials committed in application.yaml; configure all secrets externally.','Resolve the duplicate “/” page mappings and verify the intended root route.','Add authentication and authorization for UI and adjustment endpoints, including suitable CSRF protection if cookie-based sessions are used.','Confirm the deployed database name, schema, table, columns, composite key, grants, and network route with the DBA.','Define operational health probes and explicitly configure Actuator exposure and access controls.','Decide whether scheduled adjustment should be truly recurring; current implementation only runs when the POST endpoint is invoked.','Review transfer failure handling: a DC adjustment may succeed before the corresponding CA adjustment fails, leaving a partial movement that needs reconciliation.','Set log retention, monitoring, alert routing, TLS termination, firewall policy, and a rollback owner.','Consider disabling SQL logging in production and add migration/change management outside Hibernate auto-DDL.']:
 bullet(x)
para('The deployment examples are a baseline for a single Linux server. Final sizing, Oracle connectivity, secrets storage, proxy configuration, and access-control integration depend on the organization’s infrastructure and security standards.')

doc.add_heading('Appendix A. Source map',level=1)
table(['Area','Repository path'],[['Build and dependencies','pom.xml'],['Runtime configuration','src/main/resources/application.yaml'],['Application bootstrap','src/main/java/com/_apparel/ecomm_inventory_movement/EcommInventoryMovementApplication.java'],['HTTP routes and validation','src/main/java/com/_apparel/ecomm_inventory_movement/controller/InventoryController.java'],['Oracle model and repository','src/main/java/com/_apparel/ecomm_inventory_movement/entity/InventoryLi.java; repository/InventoryRepository.java'],['Inventory API calls','src/main/java/com/_apparel/ecomm_inventory_movement/service/InventoryAdjustmentService.java'],['Views','src/main/resources/templates/menu/']])

# keep table rows together where possible; set document properties
for t in doc.tables:
 for row in t.rows:
  trPr=row._tr.get_or_add_trPr(); cant=OxmlElement('w:cantSplit'); trPr.append(cant)
doc.core_properties.title='Ecomm Inventory Movement Technical Documentation'
doc.core_properties.subject='Application setup and Linux server deployment'
doc.core_properties.author='Engineering'
doc.save(out)
print(out.resolve())
